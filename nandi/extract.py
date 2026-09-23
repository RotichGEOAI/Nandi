from pathlib import Path
from .models import DocumentResult

def _plain_text(data):
    return data.decode("utf-8", errors="replace")

def extract_bytes(data: bytes, name: str, source_url: str, content_type="") -> DocumentResult:
    result = DocumentResult(name, source_url, content_type)
    suffix = Path(name.split("?")[0]).suffix.lower()
    try:
        if suffix == ".pdf" or "pdf" in content_type.lower():
            from pypdf import PdfReader
            reader = PdfReader(__import__("io").BytesIO(data))
            result.text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx" or "word" in content_type.lower():
            from docx import Document
            doc = Document(__import__("io").BytesIO(data))
            result.text = "\n".join(p.text for p in doc.paragraphs)
            result.tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
        elif "html" in content_type.lower() or suffix in {".html", ".htm"}:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(data, "html.parser")
            result.text = soup.get_text("\n", strip=True)
            result.tables = [[[cell.get_text(" ", strip=True) for cell in row.find_all(["td", "th"])]
                              for row in table.find_all("tr") if row.find_all(["td", "th"])]
                             for table in soup.find_all("table")]
        elif suffix == ".csv" or "csv" in content_type.lower():
            import csv
            from io import StringIO
            rows = list(csv.reader(StringIO(_plain_text(data))))
            result.tables = [rows] if rows else []
            result.text = "\n".join(" | ".join(row) for row in rows)
        else:
            result.text = _plain_text(data)
        if not result.text.strip() and not result.tables:
            result.status = "needs_review"
            result.review_flags.append("no_extractable_text_or_tables")
        else:
            result.status = "extracted"
    except Exception as exc:
        result.status, result.error = "error", str(exc)
        result.review_flags.append("extraction_failed")
    return result
